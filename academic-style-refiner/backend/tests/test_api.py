import io
import json

from docx import Document
from fastapi.testclient import TestClient

from app.main import app

DOC = (
    "# Introduction\n\n"
    "It is important to note that the method plays a crucial role in outcomes (Smith, 2020). "
    "Furthermore, results improved by 12%.\n\n"
    "In order to test this, we utilized 3 datasets [1].\n\n"
    "References\n\nSmith, J. (2020). Paper."
)


def parse_sse(body: str):
    events = []
    for chunk in body.strip().split("\n\n"):
        lines = dict(line.split(": ", 1) for line in chunk.split("\n") if not line.startswith(":"))
        if lines:
            events.append((lines["event"], json.loads(lines["data"])))
    return events


def test_full_job_flow():
    with TestClient(app) as client:
        r = client.post("/api/jobs", json={"text": DOC, "intensity": "strong"})
        assert r.status_code == 201
        job = r.json()
        kinds = [b["kind"] for b in job["blocks"]]
        assert kinds == ["heading", "paragraph", "paragraph", "heading", "reference"]

        with client.stream("GET", f"/api/jobs/{job['job_id']}/events") as s:
            events = parse_sse(s.read().decode())
        names = [e for e, _ in events]
        assert names[-1] == "done" and names.count("block") == 2
        done = events[-1][1]
        assert done["status"] == "done"
        assert done["metrics_after"]["words"] > 0

        snap = client.get(f"/api/jobs/{job['job_id']}").json()
        out = snap["outputs"]["1"]
        assert "(Smith, 2020)" in out and "12%" in out
        assert "It is important to note" not in out
        assert client.delete(f"/api/jobs/{job['job_id']}").status_code == 204
        assert client.get(f"/api/jobs/{job['job_id']}").status_code == 404


def test_word_limit_enforced():
    with TestClient(app) as client:
        r = client.post("/api/analyze", json={"text": "word " * 10_001})
        assert r.status_code == 413


def test_paragraph_refine():
    with TestClient(app) as client:
        r = client.post("/api/refine/paragraph", json={"text": "We utilized 3 datasets [1].", "intensity": "light"})
        assert r.status_code == 200
        assert r.json()["text"] == "We use 3 datasets [1]."


def test_extract_docx_and_exports():
    doc = Document()
    doc.add_heading("Intro", level=1)
    doc.add_paragraph("A paragraph with “curly quotes” — and an em dash.")
    doc.add_paragraph("item one", style="List Bullet")
    doc.add_paragraph("item two", style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    with TestClient(app) as client:
        r = client.post("/api/extract", files={"file": ("t.docx", buf.getvalue())})
        assert r.status_code == 200
        assert (
            r.json()["text"] == "# Intro\n\nA paragraph with “curly quotes” — and an em dash.\n\n- item one\n- item two"
        )
        assert client.post("/api/extract", files={"file": ("t.pdf", b"%PDF")}).status_code == 415

        blocks = [
            {"kind": "heading", "text": "# Intro"},
            {"kind": "paragraph", "text": "Body with <angle> & “quotes” — αβγ."},
            {"kind": "list", "text": "- a\n- b"},
        ]
        for fmt, magic in (("txt", b"# Intro"), ("docx", b"PK"), ("pdf", b"%PDF")):
            r = client.post("/api/export", json={"format": fmt, "title": "My Report", "blocks": blocks})
            assert r.status_code == 200, r.text
            assert r.content.startswith(magic)
            assert 'filename="My-Report.' in r.headers["content-disposition"]


def test_docx_export_drops_heading_that_repeats_title():
    blocks = [{"kind": "heading", "text": "# My Report"}, {"kind": "paragraph", "text": "Body."}]
    with TestClient(app) as client:
        r = client.post("/api/export", json={"format": "docx", "title": "My Report", "blocks": blocks})
    texts = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert texts == ["My Report", "Body."]
