"""Reading uploaded files and writing exports (.txt, .docx, .pdf), all in memory."""

import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .chunker import BlockKind


class UnsupportedFile(Exception):
    pass


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".txt", ".md"):
        for encoding in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise UnsupportedFile("Could not decode the text file; please save it as UTF-8.")
    if suffix == ".docx":
        try:
            doc = Document(io.BytesIO(data))
        except Exception as e:
            raise UnsupportedFile("The .docx file could not be read.") from e
        parts: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            if style.startswith("heading") or style == "title":
                level = re.search(r"\d", style)
                parts.append("#" * (int(level.group()) if level else 1) + " " + text)
            elif "list" in style:
                parts.append("- " + text)
            else:
                parts.append(text)
        # Consecutive list items belong to one block; everything else is its own block.
        out: list[str] = []
        for part in parts:
            if out and part.startswith("- ") and out[-1].split("\n")[-1].startswith("- "):
                out[-1] += "\n" + part
            else:
                out.append(part)
        return "\n\n".join(out)
    raise UnsupportedFile("Only .txt, .md and .docx files are supported.")


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

ExportBlock = tuple[BlockKind, str]


def _heading(text: str) -> tuple[int, str]:
    match = re.match(r"^(#{1,6})\s+(.*)$", text)
    if match:
        return len(match.group(1)), match.group(2)
    return 2, text


def _without_duplicate_title(title: str, blocks: list[ExportBlock]) -> list[ExportBlock]:
    """Drop a leading heading that just repeats the document title."""
    if title and blocks and blocks[0][0] == "heading" and _heading(blocks[0][1])[1].strip() == title.strip():
        return blocks[1:]
    return blocks


def to_txt(blocks: list[ExportBlock]) -> bytes:
    return ("\n\n".join(text for _, text in blocks) + "\n").encode("utf-8")


def to_docx(title: str, blocks: list[ExportBlock]) -> bytes:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    if title:
        doc.add_heading(title, level=0)
        blocks = _without_duplicate_title(title, blocks)
    for kind, text in blocks:
        if kind == "heading":
            level, body = _heading(text)
            doc.add_heading(body, level=min(level, 4))
        elif kind == "list":
            for line in text.split("\n"):
                doc.add_paragraph(line).paragraph_format.left_indent = Pt(18)
        else:
            p = doc.add_paragraph(text)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.5
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_FONT_CANDIDATES = [
    (
        "DejaVuSerif",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
    ),
    (
        "LiberationSerif",
        "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf",
    ),
]


def _pdf_fonts() -> tuple[str, str]:
    """Prefer a Unicode TrueType font when installed; fall back to the built-in Times."""
    for name, regular, bold in _FONT_CANDIDATES:
        if Path(regular).exists() and Path(bold).exists():
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, regular))
                pdfmetrics.registerFont(TTFont(f"{name}-Bold", bold))
            return name, f"{name}-Bold"
    return "Times-Roman", "Times-Bold"


def to_pdf(title: str, blocks: list[ExportBlock]) -> bytes:
    regular, bold = _pdf_fonts()
    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body", parent=base["BodyText"], fontName=regular, fontSize=11.5, leading=17, alignment=TA_JUSTIFY, spaceAfter=8
    )
    item = ParagraphStyle("Item", parent=body, leftIndent=18, alignment=0, spaceAfter=2)
    title_style = ParagraphStyle("DocTitle", parent=base["Title"], fontName=bold, fontSize=18, spaceAfter=16)
    heading_styles = {
        level: ParagraphStyle(
            f"H{level}", parent=base["Heading2"], fontName=bold, fontSize=size, spaceBefore=10, spaceAfter=6
        )
        for level, size in ((1, 16), (2, 14), (3, 12.5), (4, 12))
    }

    story = []
    if title:
        story.append(Paragraph(escape(title), title_style))
        blocks = _without_duplicate_title(title, blocks)
    for kind, text in blocks:
        if kind == "heading":
            level, heading = _heading(text)
            story.append(Paragraph(escape(heading), heading_styles[min(level, 4)]))
        elif kind == "list":
            for line in text.split("\n"):
                story.append(Paragraph(escape(line), item))
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(escape(text), body))

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        topMargin=2.5 * cm,
        bottomMargin=2.5 * cm,
        title=title or "Refined document",
    ).build(story)
    return buf.getvalue()
